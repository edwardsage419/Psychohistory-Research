#!/usr/bin/env python3
"""R091 source-only audit of fixed original Jonsson experimental workbook.

Does not output individual or group data, actions, behavioral outcome summaries,
or any source cell values other than author worksheet names and safe header tags.
"""
from __future__ import annotations
import argparse,hashlib,io,json,pathlib,re,urllib.request,urllib.error
from collections import Counter
from openpyxl import load_workbook

REPO="markusrobertjonsson/condcoop"
COMMIT="2a23fc3221e636c0119b1d1f1a9c4945bca83e85"
SRCFILE="Data_Main.xlsx"
BLOB="499787e08242c44c32a469b5d72b206ccf2c73f0"
BYTES=571844
SOURCE=f"https://raw.githubusercontent.com/{REPO}/{COMMIT}/{SRCFILE}"
EXPECTED_SHEETS=("DiffTreatments","DiffProb")
HEADER_WORDS=re.compile(r"group|round|session|treat|id|country|contrib|condition|player|subject|period|participant|number|decision|risk|prob|sweden|philippine|time|date|block|type|sample",re.IGNORECASE)
HEADER_FORMAT=re.compile(r"^[A-Za-z_][A-Za-z_0-9 .\-\(\)/]{0,59}$")

def git_blob(b):
 if not isinstance(b,bytes):raise ValueError("R091_SOURCE_GIT_HASH_REQUIRES_BYTES")
 return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

def source_identity(b,expected_blob=BLOB,expected_bytes=BYTES):
 if len(b)!=expected_bytes or git_blob(b)!=expected_blob:
  raise ValueError("R091_PINNED_AUTHOR_ORIGINAL_XLSX_SOURCE_HASH_SIZE_MISMATCH")
 return {"source_file":SRCFILE,"original_fixed_commit":COMMIT,
         "source_git_blob_sha1":git_blob(b),"full_source_SHA256":hashlib.sha256(b).hexdigest(),
         "original_source_bytes":len(b)}

def permitted_header(cell):
 if not isinstance(cell,str):return None
 s=cell.strip()
 if not (HEADER_FORMAT.fullmatch(s) and HEADER_WORDS.search(s)):return None
 return s

def shape_only(workbook):
 if tuple(workbook.sheetnames)!=EXPECTED_SHEETS:
  raise ValueError("R091_ORIGINAL_AUTHOR_WORKSHEETS_UNEXPECTED")
 result={}
 for name in EXPECTED_SHEETS:
  ws=workbook[name]
  # Do not store any raw cells. Safely distinguish numeric layout and headers.
  first_two=[]
  for row in ws.iter_rows(min_row=1,max_row=min(2,ws.max_row),values_only=True):
   first_two.append([{"col":k+1,"safe":permitted_header(v)} for k,v in enumerate(row)
                     if permitted_header(v) is not None])
  coltypes=[]
  for c in range(1,ws.max_column+1):
   counts=Counter()
   for values in ws.iter_rows(min_col=c,max_col=c,values_only=False):
    cell=values[0];v=cell.value
    if v is None:counts["empty"]+=1
    elif cell.data_type=="f":counts["formula"]+=1
    elif isinstance(v,bool):counts["boolean"]+=1
    elif isinstance(v,(int,float)):counts["numeric"]+=1
    elif isinstance(v,str):counts["string"]+=1
    else:counts["other_typed"]+=1
   coltypes.append({"column_index":c,"cell_types":{k:counts[k] for k in
    ("numeric","string","boolean","formula","empty","other_typed")}})
  result[name]={"sheet_rows_including_headers":ws.max_row,"sheet_columns":ws.max_column,
    "safe_column_name_candidates_in_first_two_rows":first_two,
    "strict_anonymous_per_column_cell_type_counts":coltypes}
 return result

def download():
 req=urllib.request.Request(SOURCE,headers={"User-Agent":"Psychohistory-R091-author-original-Git-SHA-XLSX-audit/1.0"})
 with urllib.request.urlopen(req,timeout=60) as res:
  if res.geturl()!=SOURCE:raise ValueError("R091_FROZEN_AUTHOR_SOURCE_REDIRECT")
  payload=res.read(BYTES+1)
 identity=source_identity(payload)
 workbook=load_workbook(io.BytesIO(payload),read_only=True,data_only=False)
 try:shapes=shape_only(workbook)
 finally:workbook.close()
 return {"research_id":"R091","phase":"A_FIXED_FIRST_PARTY_XLSX_SOURCE_STRUCTURE_ONLY",
   "source_identity":identity,"first_party_original_two_sheets":shapes,
   "original_paper_DOI":"10.1371/journal.pone.0318891",
   "paper_reported_original_human_subjects":884,
   "author_published_workbook_sheets_not_assumed_disjoint":True,
   "original_person_group_or_behavior_values_exported":False,
   "person_or_group_identifiers_or_hashes_exported":False,
   "actual_human_cooperation_scored":False,
   "R091_new_real_E2_group_history_prediction":"BLOCKED_PENDING_SEPARATELY_FROZEN_ACTOR_GROUP_ROUND_ELIGIBILITY",
   "PESEN_v0.1":"FROZEN"}

def main():
 ap=argparse.ArgumentParser();ap.add_argument("--output",type=pathlib.Path,required=True)
 a=ap.parse_args()
 try:r=download()
 except (ValueError,TypeError,OSError,urllib.error.HTTPError,urllib.error.URLError) as e:
  msg=str(e)[:100] if isinstance(e,ValueError) and str(e).startswith("R091_") else "R091_XLSX_SOURCE_OR_WORKBOOK_LAYOUT_FAIL_CLOSED"
  r={"research_id":"R091","phase":"A_ORIGINAL_XLSX_SOURCE_FAIL_CLOSED",
     "safe_failure":msg,"actual_human_cooperation_scored":False,"PESEN_v0.1":"FROZEN"}
 a.output.parent.mkdir(parents=True,exist_ok=True)
 a.output.write_text(json.dumps(r,sort_keys=True,indent=2)+"\n",encoding="utf-8")
 brief=r.get("first_party_original_two_sheets",{})
 print(json.dumps({"phase":r["phase"],"safe_failure":r.get("safe_failure"),
  "identity":r.get("source_identity"),
  "sheets":{k:{"rows":v["sheet_rows_including_headers"],"columns":v["sheet_columns"],
    "header_names_first_two":v["safe_column_name_candidates_in_first_two_rows"]}
    for k,v in brief.items()},"original_outcome_scored":False},sort_keys=True))
if __name__=="__main__":main()

#!/usr/bin/env python3
"""R091 fixed source categorical labels only, declassified bounded categories.

No actor/group/period identifiers or individual contributions are processed.
"""
from __future__ import annotations
import argparse,collections,io,json,pathlib,re,urllib.request,urllib.error
from collections import Counter
from openpyxl import load_workbook
from r091_author_2025_disaster_xlsx_source_only_gate import (
 SOURCE,BYTES,EXPECTED_SHEETS,source_identity)
EXPECTED_SHA256="04d2ed5a59885cde2a1740f33b295a86bd374b80cd811d8ff84c88cded68a023"
LEGAL=re.compile(r"^[A-Za-z0-9._% -]{1,30}$")

def category_label(value):
 if value is None:return "MISSING_LABEL"
 s=str(value).strip()
 if not LEGAL.fullmatch(s):return "REDACTED_UNRECOGNIZED_LABEL"
 return s

def metadata_codes(rows,expected=15200):
 if len(rows)!=expected:raise ValueError("R091B2_ORIGINAL_MAIN_SHEET_ROWS_NOT15200")
 cats={"Country":Counter(),"Treatment":Counter()}
 for r in rows:
  if len(r)!=2:raise ValueError("R091B2_EXPECT_EXACT_CATEGORICAL_SOURCE_COLUMNS")
  cats["Country"][category_label(r[0])]+=1
  cats["Treatment"][category_label(r[1])]+=1
 for col,d in cats.items():
  if len(d)>32:raise ValueError("R091B2_TOO_MANY_CATEGORICAL_VALUES")
 return {"original_main_actor_round_rows":len(rows),
  "Country_safe_original_metadata_label_frequency":{k:v for k,v in sorted(cats["Country"].items())},
  "Treatment_safe_original_metadata_label_frequency":{k:v for k,v in sorted(cats["Treatment"].items())},
  "contribution_values_examined":False,
  "actor_group_round_identifiers_exported":False}

def run_source():
 with urllib.request.urlopen(urllib.request.Request(SOURCE,headers={
  "User-Agent":"Psychohistory-R091-categorical-source-schema-bounded-only/1.0"}),timeout=55) as r:
  if r.geturl()!=SOURCE:raise ValueError("R091B2_UNPINNED_ORIGINAL_REDIRECT")
  raw=r.read(BYTES+1)
 ident=source_identity(raw)
 if ident["full_source_SHA256"]!=EXPECTED_SHA256:raise ValueError("R091B2_SOURCE_SHA256_MISMATCH")
 w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 try:
  if tuple(w.sheetnames)!=EXPECTED_SHEETS:raise ValueError("R091B2_ORIGINAL_SHEETS_CHANGED")
  ws=w["DiffTreatments"]
  if ws.max_row!=15201 or ws.max_column!=6:
   raise ValueError("R091B2_SOURCE_15201X6_CHANGED")
  rows=[x for x in ws.iter_rows(min_row=2,min_col=5,max_col=6,values_only=True)]
  audit=metadata_codes(rows)
 finally:w.close()
 return {"research_id":"R091","phase":"B2_ORIGINAL_AUTHOR_CATEGORICAL_LABELS_ONLY",
  "source_identity":ident,"safe_original_treatment_and_country_labels":audit,
  "first_source_country_treatment_mismatch_failure_retained":True,
  "real_four_person_group_scores_computed":False,
  "PESEN_v0.1":"FROZEN"}

def main():
 ap=argparse.ArgumentParser();ap.add_argument("--output",type=pathlib.Path,required=True);a=ap.parse_args()
 try:out=run_source()
 except (ValueError,TypeError,OSError,urllib.error.URLError,urllib.error.HTTPError) as e:
  reason=str(e)[:100] if isinstance(e,ValueError) and str(e).startswith("R091B2_") else "R091B2_AUTHOR_SOURCE_UNREADABLE"
  out={"research_id":"R091","phase":"B2_SOURCE_CATEGORY_FAIL_CLOSED",
   "safe_failure":reason,"real_four_person_group_scores_computed":False,"PESEN_v0.1":"FROZEN"}
 a.output.parent.mkdir(parents=True,exist_ok=True)
 a.output.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
 print(json.dumps(out,sort_keys=True))
if __name__=="__main__":main()

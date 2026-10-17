#!/usr/bin/env python3
"""R091 independent first-party Jonsson disaster 4-person 20-round LOOG E2.

Frozen prior to any original behavioral evaluation. All source IDs/row-level
contributions strictly in-memory. Output anonymous country×treatment macro only.
"""
from __future__ import annotations
import argparse,collections,io,json,math,pathlib,urllib.request,urllib.error
from collections import Counter,defaultdict
from openpyxl import load_workbook
from r091_author_2025_disaster_xlsx_source_only_gate import (
 SOURCE,BYTES,source_identity,EXPECTED_SHEETS)
from r091_phase_b_main_sheet_group_time_roster_only import (
 qualify_rows,country_tag,condition_tag,time_round)

SOURCE_SHA256="04d2ed5a59885cde2a1740f33b295a86bd374b80cd811d8ff84c88cded68a023"
PROTOCOL="protocols/R091_PHASE_C_190_AUTHOR_4PLAYER_GROUP_COMPLETE_LOOG_M0_M1_M2_BEHAVIOR_SCORING_FREEZE_2026-10-17.md"
ARMS={
 "Philippines":{"Control":22,"10P":22,"40P":23,"Impact":22,"Level":22},
 "Sweden":{"Control":18,"10P":18,"40P":16,"Impact":15,"Level":12}}
MODELS=("M0_OTHER_SAME_COUNTRY_TREATMENT_ROUND","M1_OWN_GROUP_PREVIOUS_ROUND","M2_FIXED_HALF_BLEND")
COUNTRIES=("Philippines","Sweden")
TREATMENTS=("Control","10P","40P","Impact","Level")

def contribution(x):
 if isinstance(x,bool) or not isinstance(x,(float,int)):
  raise ValueError("R091C_ORIGINAL_NONNUMERIC_MISSING_OR_FORMULA_CONTRIBUTION")
 if not math.isfinite(x) or int(x)!=x or not 0<=x<=20:
  raise ValueError("R091C_ORIGINAL_CONTRIBUTION_NOT_INTEGER_IN_0_TO_20")
 return int(x)

def source_groups(rows,expected_total=15200,expected_arm_counts=ARMS):
 # The source roster qualification requires 4 participants and 20 rounds,
 # with anonymous ID structure verified separately and unchanged.
 expected_g=sum(sum(d.values()) for d in expected_arm_counts.values())
 qualify_rows(rows,expected_total,expected_g)
 actors=defaultdict(lambda:defaultdict(dict))
 grouptags=defaultdict(set)
 for row in rows:
  actor,g,amount,p,c,t=row[:6]
  country=country_tag(c);condition=condition_tag(t)
  if country not in expected_arm_counts or condition not in expected_arm_counts[country]:
   raise ValueError("R091C_ORIGINAL_SOURCE_COUNTRY_CONDITION_CHANGED")
  groupkey=(country,condition,str(g).strip())
  when=time_round(p)
  uid=str(actor).strip()
  if not uid:raise ValueError("R091C_MISSING_ORIGINAL_ACTOR_ID")
  value=contribution(amount)
  if uid in actors[groupkey][when]:
   raise ValueError("R091C_SOURCE_DUPLICATE_PERSON_GROUP_ROUND")
  actors[groupkey][when][uid]=value
  grouptags[(country,condition)].add(groupkey)
 groups=defaultdict(dict)
 for key,byround in actors.items():
  if set(byround)!=set(range(1,21)):
   raise ValueError("R091C_MISSING_20_ROUND_GROUP")
  first=frozenset(byround[1])
  if len(first)!=4 or any(frozenset(byround[t])!=first for t in range(1,21)):
   raise ValueError("R091C_FOUR_PLAYER_ORIGINAL_GROUP_MEMBERSHIP_CHANGED")
  totals=[sum(byround[t].values())/80.0 for t in range(1,21)]
  if any(not 0<=z<=1 for z in totals):
   raise ValueError("R091C_GROUP_COOPERATION_OUTSIDE_NORMALIZED_RANGE")
  country,condition,_=key
  groups[(country,condition)][key]=totals
 for country,conds in expected_arm_counts.items():
  for cond,n in conds.items():
   if len(groups[(country,cond)])!=n:
    raise ValueError("R091C_ORIGINAL_SOURCE_10_ARM_190_GROUP_COUNTS_CHANGED")
 return groups

def mean(seq):
 if not seq:raise ValueError("R091C_EMPTY_GROUP_SCORE_DENOMINATOR")
 return sum(seq)/len(seq)

def evaluate_group_loo(groups,expected_arm_counts=ARMS,expected_targets=3610):
 per_arm={}
 global_rows=[]
 country_rows=defaultdict(list)
 treatment_rows=defaultdict(list)
 for country in COUNTRIES:
  for condition in TREATMENTS:
   members=groups[(country,condition)]
   want=expected_arm_counts[country][condition]
   if len(members)!=want or len(members)<2:
    raise ValueError("R091C_NO_TREATMENT_COUNTRY_OUTGROUP_TRAINING")
   dataset=list(members.items())
   records=[];good=bad=zero=0
   byround=defaultdict(list)
   for group_id,periods in dataset:
    if len(periods)!=20 or any(not isinstance(x,(int,float)) or not math.isfinite(x) or not 0<=x<=1 for x in periods):
     raise ValueError("R091C_INVALID_ORIGINAL_NORMALIZED_GROUP_RECORD")
    other_groups=[v for oid,v in dataset if oid!=group_id]
    if len(other_groups)!=want-1:raise ValueError("R091C_HELDOUT_SOURCE_GROUP_LEAKAGE")
    gains=[]
    for idx in range(1,20):
     actual=periods[idx]
     p0=mean([v[idx] for v in other_groups])
     p1=periods[idx-1]
     p2=0.5*(p0+p1)
     e0=(p0-actual)**2;e1=(p1-actual)**2;e2=(p2-actual)**2
     record=(e0,e1,e2,e0-e1,e0-e2)
     records.append(record)
     byround[idx+1].append(actual)
     gains.append(e0-e1)
    s=mean(gains)
    good+=int(s>0);bad+=int(s<0);zero+=int(s==0)
   if len(records)!=want*19 or good+bad+zero!=want:
    raise ValueError("R091C_ORIGINAL_HELDOUT_GROUP_ROUND_COUNT_NOT_COMPLETE")
   pooled=list(records)
   global_rows.extend(pooled);country_rows[country].extend(pooled);treatment_rows[condition].extend(pooled)
   per_arm[f"{country}|{condition}"]={
    "country":country,"original_experiment_treatment":condition,
    "original_fourperson_groups":want,"original_actor_round_contribution_slots":want*4*20,
    "heldout_next_round_group_predictions":len(records),
    "M0_other_groups_same_stratum_MSE":mean([z[0] for z in records]),
    "M1_same_group_last_round_MSE":mean([z[1] for z in records]),
    "M2_fixed_halfblend_MSE":mean([z[2] for z in records]),
    "paired_gain_M0_minus_M1":mean([z[3] for z in records]),
    "paired_gain_M0_minus_M2":mean([z[4] for z in records]),
    "full_groups_M1_gain_positive":good,
    "full_groups_M1_gain_negative":bad,
    "full_groups_M1_gain_zero":zero,
    "author_heldout_group_next_round_cooperation_mean_by_round_2_to_20":[mean(byround[i]) for i in range(2,21)],
    "no_group_or_person_original_source_ids_in_output":True}
 if len(global_rows)!=expected_targets:
  raise ValueError("R091C_HELDOUT_GROUP_TARGETS_NOT3610")
 def aggregate(rows):
  return {"original_heldout_next_round_group_targets":len(rows),
   "M0_MSE":mean([z[0] for z in rows]),
   "M1_MSE":mean([z[1] for z in rows]),
   "M2_MSE":mean([z[2] for z in rows]),
   "paired_gain_M0_minus_M1":mean([z[3] for z in rows]),
   "paired_gain_M0_minus_M2":mean([z[4] for z in rows])}
 overall=aggregate(global_rows)
 groups_better=sum(x["full_groups_M1_gain_positive"] for x in per_arm.values())
 groups_worse=sum(x["full_groups_M1_gain_negative"] for x in per_arm.values())
 groups_zero=sum(x["full_groups_M1_gain_zero"] for x in per_arm.values())
 if groups_better+groups_worse+groups_zero!=sum(sum(d.values()) for d in expected_arm_counts.values()):
  raise ValueError("R091C_GROUP_GAIN_SIGN_DENOMINATOR_FAILED")
 return {
  "original_fourperson_groups_heldout":sum(sum(d.values()) for d in expected_arm_counts.values()),
  "original_4player_20round_people":sum(sum(d.values()) for d in expected_arm_counts.values())*4,
  "original_4player_actor_round_contribution_slots":sum(sum(d.values()) for d in expected_arm_counts.values())*80,
  "original_independent_group_next_round_test_targets":len(global_rows),
  "pooled_M0_M1_M2":overall,
  "by_original_country":{c:aggregate(country_rows[c]) for c in COUNTRIES},
  "by_original_treatment":{t:aggregate(treatment_rows[t]) for t in TREATMENTS},
  "by_original_country_and_treatment":per_arm,
  "all_source_groups_M1_gain_positive":groups_better,
  "all_source_groups_M1_gain_negative":groups_worse,
  "all_source_groups_M1_gain_zero":groups_zero,
  "pre_frozen_hypothesis_M1_gain_positive":overall["paired_gain_M0_minus_M1"]>0,
  "pre_frozen_hypothesis_M2_gain_positive":overall["paired_gain_M0_minus_M2"]>0,
  "M1_positive_in_all_10_original_country_treatment_strata":all(v["paired_gain_M0_minus_M1"]>0 for v in per_arm.values()),
  "M2_positive_in_all_10_original_country_treatment_strata":all(v["paired_gain_M0_minus_M2"]>0 for v in per_arm.values()),
  "group_or_actor_IDs_or_hashed_source_features_in_output":False}

def real_author_experiment():
 with urllib.request.urlopen(urllib.request.Request(SOURCE,headers={
  "User-Agent":"Psychohistory-R091-independent-author-original-190group-full-Git-blob/1.0"}),timeout=60) as r:
  if r.geturl()!=SOURCE:raise ValueError("R091C_AUTHOR_ORIGINAL_SOURCE_UNEXPECTED_REDIRECT")
  raw=r.read(BYTES+1)
 identity=source_identity(raw)
 if identity["full_source_SHA256"]!=SOURCE_SHA256:
  raise ValueError("R091C_AUTHOR_FULL_ORIGINAL_SOURCE_SHA256_CHANGED")
 w=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
 try:
  if tuple(w.sheetnames)!=EXPECTED_SHEETS:raise ValueError("R091C_ORIGINAL_SOURCE_TWO_SHEETS_CHANGED")
  ws=w["DiffTreatments"]
  if ws.max_column!=6 or ws.max_row!=15201:
   raise ValueError("R091C_ORIGINAL_AUTHOR_MAIN_WORKSHEET_DIMENSIONS_CHANGED")
  rows=[r for r in ws.iter_rows(min_row=2,max_col=6,values_only=True)]
 finally:w.close()
 groups=source_groups(rows)
 result=evaluate_group_loo(groups)
 return {
  "research_id":"R091",
  "status":"R091_REAL_ORIGINAL_AUTHOR_190_FOURPERSON_GROUP_FULL_HOLDOUT_E2_EXECUTED",
  "evidence_grade":"E2_RETROSPECTIVE_INDEPENDENT_ORIGINAL_AUTHOR_4PLAYER_DATA_M0_M1_M2_TRANSFER",
  "original_paper_DOI":"10.1371/journal.pone.0318891",
  "original_author_fixed_Git_commit":"2a23fc3221e636c0119b1d1f1a9c4945bca83e85",
  "original_author_Git_blob":identity["source_git_blob_sha1"],
  "original_author_source_SHA256":identity["full_source_SHA256"],
  "original_behavior_scoring_frozen_protocol":PROTOCOL,
  "original_main_sheet_only":"DiffTreatments",
  "original_second_sheet_not_appended_to_avoid_double_count":True,
  "original_full_190_group_20round_prediction_results":result,
  "original_paper_reported_884_participants_not_equated_to_main_sheet_760":True,
  "new_independent_human_recruitment_or_prospective_E3":False,
  "causal_disaster_risk_group_cooperation_mechanism":"NOT_IDENTIFIED_BY_THIS_SCORE",
  "source_original_automatic_decision_flags_available":False,
  "individual_original_source_contribution_or_group_IDs_exported":False,
  "PESEN_v0.1":"FROZEN"}

def main():
 p=argparse.ArgumentParser();p.add_argument("--output",type=pathlib.Path,required=True);a=p.parse_args()
 try:r=real_author_experiment()
 except (ValueError,OSError,TypeError,KeyError,urllib.error.URLError,urllib.error.HTTPError) as e:
  safe=str(e)[:100] if isinstance(e,ValueError) and str(e).startswith("R091C_") else "R091C_AUTHOR_SOURCE_OR_ORIGINAL_BEHAVIORAL_EVALUATION_FAIL_CLOSED"
  r={"research_id":"R091","status":"R091C_ORIGINAL_REAL_BEHAVIORAL_SOURCE_ELIGIBILITY_FAIL_CLOSED",
     "safe_failure":safe,"actual_real_actor_group_cooperation_scores_exported":0,"PESEN_v0.1":"FROZEN"}
 a.output.parent.mkdir(parents=True,exist_ok=True)
 a.output.write_text(json.dumps(r,sort_keys=True,indent=2)+"\n",encoding="utf-8")
 scores=r.get("original_full_190_group_20round_prediction_results",{})
 print(json.dumps({"status":r["status"],"safe_failure":r.get("safe_failure"),
  "results":scores.get("pooled_M0_M1_M2"),"country":scores.get("by_original_country"),
  "treatments":scores.get("by_original_treatment"),
  "positive_groups":scores.get("all_source_groups_M1_gain_positive"),
  "negative_groups":scores.get("all_source_groups_M1_gain_negative")},sort_keys=True))
 if r["status"]!="R091_REAL_ORIGINAL_AUTHOR_190_FOURPERSON_GROUP_FULL_HOLDOUT_E2_EXECUTED":
  raise RuntimeError("R091C_ORIGINAL_REAL_GIT_SOURCE_BEHAVIOR_FAILED")
if __name__=="__main__":main()

#!/usr/bin/env python3
"""R091 synthetic-only author XLSX blob and source-layout privacy tests."""
import io,pathlib,sys,unittest
from openpyxl import Workbook,load_workbook
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parent))
from r091_author_2025_disaster_xlsx_source_only_gate import (
 git_blob,source_identity,permitted_header,shape_only,EXPECTED_SHEETS,
 BYTES,BLOB,COMMIT,SOURCE)
def fixture():
 w=Workbook();w.active.title="DiffTreatments"
 w.active.append(["group_id","round","contribution","participant_code"])
 w.active.append(["PRIVATE-PERSON-GROUP999",1,19,"PRIVATE-ACTOR-666"])
 w.create_sheet("DiffProb")
 w["DiffProb"].append(["group_id","round","contribution"])
 w["DiffProb"].append(["PRIVATE-GROUP-SECRET",2,20])
 return w
class SourceRules(unittest.TestCase):
 def test_sha_original_blob_identity(self):self.assertEqual(len(BLOB),40)
 def test_original_commit_frozen(self):self.assertEqual(len(COMMIT),40)
 def test_fixed_author_original_file_bytes(self):self.assertEqual(BYTES,571844)
 def test_repo_commit_url(self):self.assertIn(COMMIT,SOURCE)
 def test_known_empty_blob_sha(self):self.assertEqual(git_blob(b""),"e69de29bb2d1d6434b8b29ae775ad8c2e48c5391")
 def test_changed_source_sha(self):self.assertNotEqual(git_blob(b"A"),git_blob(b"B"))
 def test_invalid_source_sha_fails(self):
  with self.assertRaises(ValueError):source_identity(b"not author original")
 def test_sha_check_synthetic_exact(self):
  self.assertEqual(source_identity(b"test",git_blob(b"test"),4)["original_source_bytes"],4)
 def test_named_expected_sheets(self):self.assertEqual(EXPECTED_SHEETS,("DiffTreatments","DiffProb"))
 def test_header_allowed_group(self):self.assertEqual(permitted_header("group_id"),"group_id")
 def test_header_allowed_contrib(self):self.assertEqual(permitted_header("contribution"),"contribution")
 def test_header_allow_round(self):self.assertEqual(permitted_header("Round"),"Round")
 def test_header_actor_private_value_denied(self):self.assertIsNone(permitted_header("PRIVATE-ACTOR-666"))
 def test_header_numeric_denied(self):self.assertIsNone(permitted_header(12))
 def test_header_raw_free_text_denied(self):self.assertIsNone(permitted_header("actor says I'll never cooperate"))
 def test_layout_type_totals(self):
  v=shape_only(fixture());self.assertEqual(v["DiffTreatments"]["sheet_rows_including_headers"],2)
  self.assertEqual(v["DiffTreatments"]["sheet_columns"],4)
 def test_only_safe_header_not_raw_cells(self):
  v=str(shape_only(fixture()))
  self.assertNotIn("PRIVATE-GROUP-SECRET",v)
  self.assertNotIn("PRIVATE-ACTOR-666",v)
  self.assertNotIn("PRIVATE-PERSON-GROUP999",v)
 def test_remove_sheet_fails(self):
  w=fixture();del w["DiffProb"]
  with self.assertRaises(ValueError):shape_only(w)
 def test_sheet_swap_order_fails(self):
  w=fixture();w._sheets.reverse()
  with self.assertRaises(ValueError):shape_only(w)
 def test_empty_and_numeric_types(self):
  w=fixture();w.active["E1"]="Group_ID";w.active["E3"]=0
  v=shape_only(w)
  self.assertEqual(v["DiffTreatments"]["strict_anonymous_per_column_cell_type_counts"][4]["cell_types"]["numeric"],1)
 def test_round_data_is_not_in_header(self):
  v=shape_only(fixture())
  self.assertNotIn("19",str(v["DiffTreatments"]["safe_column_name_candidates_in_first_two_rows"]))
if __name__=="__main__":unittest.main()

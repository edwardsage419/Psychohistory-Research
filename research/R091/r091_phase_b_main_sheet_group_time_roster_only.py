#!/usr/bin/env python3
"""R091 Jonsson 2025 original XLSX phase B source group roster ONLY.

Fixed author source commit+Git blob+whole-file SHA. Never outputs actor/group
keys, participant choices, source row values, or group contribution outcomes.
"""
from __future__ import annotations
import argparse,collections,io,json,pathlib,urllib.request,urllib.error
from collections import Counter,defaultdict
from openpyxl import load_workbook
from r091_author_2025_disaster_xlsx_source_only_gate import (
 source_identity,SOURCE,BYTES,BLOB,EXPECTED_SHEETS)

EXPECTED_SHA256="04d2ed5a59885cde2a1740f33b295a86bd374b80cd811d8ff84c88cded68a023"
CONDITIONS=("Control","10P","40P","Impact","Level")
COUNTRY_ALIASES={
 "sweden":"Sweden","swe":"Sweden","se":"Sweden",
 "philippines":"Philippines","philippine":"Philippines","phi":"Philippines","ph":"Philippines","phils":"Philippines"}
CONDITION_ALIASES={k.lower():k for k in CONDITIONS}
CONDITION_ALIASES.update({"10prob":"10P","40prob":"40P","40impact":"Impact","40level":"Level"})

def source_string(value):
 if value is None or isinstance(value,bool):return ""
 return str(value).strip()

def country_tag(value):
 return COUNTRY_ALIASES.get(source_string(value).lower())

def condition_tag(value):
 return CONDITION_ALIASES.get(source_string(value).lower())

def time_round(value):
 if isinstance(value,bool):raise ValueError("R091B_SOURCE_ROUND_NOT_1_TO_20")
 if isinstance(value,(int,float)):
  if int(value)!=value or not 1<=value<=20:
   raise ValueError("R091B_SOURCE_ROUND_NOT_1_TO_20")
  return int(value)
 if isinstance(value,str) and value.strip().isascii() and value.strip().isdigit() and 1<=int(value.strip())<=20:
  return int(value.strip())
 raise ValueError("R091B_SOURCE_ROUND_NOT_1_TO_20")

def qualify_rows(rows,expected_people_rounds=15200,expected_groups=190):
 if len(rows)!=expected_people_rounds:
  raise ValueError("R091B_AUTHOR_MAIN_SHEET_ORIGINAL_ROW_NUMBER_UNEXPECTED")
 units=defaultdict(list)
 group_round=defaultdict(list)
 by_person=defaultdict(set)
 condition_groups=defaultdict(set)
 original_actor_id_empty=0
 for row in rows:
  if len(row)<6:raise ValueError("R091B_ORIGINAL_MAIN_SHEET_SIX_COLS_REQUIRED")
  actor,group,_,period,country,condition=row[:6]
  g=source_string(group)
  if not g:raise ValueError("R091B_ORIGINAL_SOURCE_GROUP_CODE_MISSING")
  c=country_tag(country);t=condition_tag(condition)
  if c is None or t is None:
   raise ValueError("R091B_AUTHOR_COUNTRY_OR_TREATMENT_CODE_UNMAPPED")
  p=time_round(period)
  actor=source_string(actor)
  original_actor_id_empty+=int(not actor)
  key=(c,t,g)
  units[key].append((p,actor))
  group_round[(key,p)].append(actor)
  if actor:by_person[(key,actor)].add(p)
  condition_groups[(c,t)].add(key)
 if len(units)!=expected_groups or len(group_round)!=expected_groups*20:
  raise ValueError("R091B_SOURCE_COMPLETE_190_GROUP_20_ROUND_UNIVERSE_FAILED")
 group_round_sizes=Counter(len(v) for v in group_round.values())
 if group_round_sizes!={4:expected_groups*20}:
  raise ValueError("R091B_SOURCE_GROUP_ROUND_NOT_FOUR_SUBJECTS")
 person_id_roster_evidence=Counter()
 for key,participants in units.items():
  if len(participants)!=80 or set(p for p,a in participants)!=set(range(1,21)):
   raise ValueError("R091B_GROUP_HAS_LESS_THAN_4x20_SOURCE_SLOTS")
  by_round=[group_round[(key,p)] for p in range(1,21)]
  stable=(all(len(set(ids))==4 and all(ids) for ids in by_round)
         and len({frozenset(ids) for ids in by_round})==1)
  person_id_roster_evidence["first_column_stable_within_original_group" if stable
                             else "first_column_not_proven_stable_as_actor_ID"]+=1
 groups_by_condition={
  c:{t:{"number_of_groups":len(condition_groups.get((c,t),set()))} for t in CONDITIONS}
  for c in ("Sweden","Philippines")}
 # Long group-period tables are anonymous; do not return group identity or label IDs.
 return {
  "verified_main_sheet_actor_round_entries":len(rows),
  "verified_4_member_20_round_source_group_count":len(units),
  "verified_complete_original_source_group_rounds":len(group_round),
  "original_source_20_rounds_per_group":True,
  "all_original_4_actor_decisions_present_per_group_round":True,
  "source_four_person_group_round_size_histogram":{str(k):v for k,v in sorted(group_round_sizes.items())},
  "source_first_column_actor_identifier_empty_cells":original_actor_id_empty,
  "roster_person_identifier_evidence_counts":dict(person_id_roster_evidence),
  "source_groups_by_country_and_treatment_safe_code_only":groups_by_condition,
  "source_unique_country_condition_arms":sum(bool(v["number_of_groups"])
    for c in groups_by_condition.values() for v in c.values()),
  "original_diffprob_sheet_not_appended_due_possible_overlap":True,
  "individual_source_behaviors_or_contributions_read_for_scoring":False,
  "source_actor_group_ids_or_hashes_exported":False}

def actual_author_roster():
 req=urllib.request.Request(SOURCE,headers={"User-Agent":"Psychohistory-R091-pinned-Jonsson-original-4player-roster-only/1.0"})
 with urllib.request.urlopen(req,timeout=60) as res:
  if res.geturl()!=SOURCE:raise ValueError("R091B_ORIGINAL_AUTHOR_SOURCE_REDIRECT")
  raw=res.read(BYTES+1)
 identity=source_identity(raw)
 if identity["full_source_SHA256"]!=EXPECTED_SHA256:
  raise ValueError("R091B_SOURCE_FULL_ORIGINAL_SHA256_MISMATCH")
 w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 try:
  if tuple(w.sheetnames)!=EXPECTED_SHEETS:
   raise ValueError("R091B_AUTHOR_TWO_SOURCE_SHEETS_CHANGED")
  ws=w["DiffTreatments"]
  if ws.max_column!=6 or ws.max_row!=15201:
   raise ValueError("R091B_MAIN_SHEET_ORIGINAL_DIMENSIONS_MISMATCH")
  # The contribution field at index 2 is not used for any R091B calculations.
  input_rows=(row for row in ws.iter_rows(min_row=2,max_col=6,values_only=True))
  report=qualify_rows(list(input_rows))
 finally:w.close()
 return {"research_id":"R091","phase":"B_ORIGINAL_2025_AUTHOR_4PERSON_GROUP_ROSTER_ONLY",
  "original_xlsx_source_identity":identity,"original_roster_eligibility":report,
  "original_behavioral_scores_computed":False,
  "independent_model_information_transfer_E2":"AWAITING_PHASE_C_FROZEN_SCORING",
  "PESEN_v0.1":"FROZEN"}

def main():
 p=argparse.ArgumentParser();p.add_argument("--output",type=pathlib.Path,required=True);a=p.parse_args()
 try:r=actual_author_roster()
 except (ValueError,OSError,TypeError,urllib.error.URLError,urllib.error.HTTPError) as e:
  safe=str(e)[:100] if isinstance(e,ValueError) and str(e).startswith("R091B_") else "R091B_SOURCE_OR_GROUP_ROSTER_FAIL_CLOSED"
  r={"research_id":"R091","phase":"B_FIRST_PARTY_GROUP_ROSTER_FAIL_CLOSED",
    "safe_failure":safe,"original_behavioral_scores_computed":False,"PESEN_v0.1":"FROZEN"}
 a.output.parent.mkdir(parents=True,exist_ok=True)
 a.output.write_text(json.dumps(r,sort_keys=True,indent=2)+"\n",encoding="utf-8")
 print(json.dumps(r,sort_keys=True))
if __name__=="__main__":main()

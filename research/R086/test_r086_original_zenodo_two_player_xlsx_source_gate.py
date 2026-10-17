#!/usr/bin/env python3
"""R086 purely synthetic source and 10-sheet Excel schema contracts."""
import io,hashlib,pathlib,sys,unittest
from openpyxl import Workbook
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parent))
from r086_original_zenodo_two_player_xlsx_source_gate import (
 SOURCES,ORIGINAL_RECORD,ORIGINAL_DOI,MAX_RAW_BYTES,exact_url,verify_url,source_summary,safe_failure,urllib)

def fake_book():
 b=io.BytesIO();w=Workbook();w.remove(w.active)
 for arm in ("FE","EI","PI","AI","MI"):
  for kind in ("Type 1","Type 2"):
   x=w.create_sheet(f"{arm}_{kind}")
   x.cell(1,1,"fake header")
   x.cell(10,22,42)
 w.save(b)
 return b.getvalue()

def fake_spec(raw):
 return {"original_author_md5":hashlib.md5(raw).hexdigest(),
         "published_number_of_humans_by_arm":{"FE":1},
         "author_first_session_rounds":20,"author_second_session_rounds":20}

class OfficialSourceContract(unittest.TestCase):
 def test_stable_record(self):self.assertEqual(ORIGINAL_RECORD,16918146)
 def test_stable_doi(self):self.assertEqual(ORIGINAL_DOI,"10.5281/zenodo.16918146")
 def test_two_separate_games(self):self.assertEqual(len(SOURCES),2)
 def test_published_linear_md5(self):
  self.assertEqual(SOURCES["two_person_linear"]["original_author_md5"],"fc71d20de62c259629ceb56559d8a69e")
 def test_published_threshold_md5(self):
  self.assertEqual(SOURCES["two_person_threshold"]["original_author_md5"],"c40171abb172da8e16b0839db98a2cf2")
 def test_published_count_is_550_linear_and_554_threshold(self):
  self.assertEqual(sum(SOURCES["two_person_linear"]["published_number_of_humans_by_arm"].values()),550)
  self.assertEqual(sum(SOURCES["two_person_threshold"]["published_number_of_humans_by_arm"].values()),554)
 def test_linear_has_session_length_22_21(self):
  self.assertEqual([SOURCES["two_person_linear"]["author_first_session_rounds"],
                    SOURCES["two_person_linear"]["author_second_session_rounds"]],[22,21])
 def test_threshold_has_20_20(self):
  self.assertEqual([SOURCES["two_person_threshold"]["author_first_session_rounds"],
                    SOURCES["two_person_threshold"]["author_second_session_rounds"]],[20,20])
 def test_two_original_exact_urls_allowed(self):
  self.assertTrue(all(verify_url(exact_url(v["filename"])) for v in SOURCES.values()))
 def test_unknown_excel_rejected(self):
  with self.assertRaises(ValueError):exact_url("human_subject_information.xlsx")
 def test_wrong_record_denied(self):
  with self.assertRaises(ValueError):verify_url("https://zenodo.org/records/1/files/LinearPGG_ExperimentalData.xlsx")
 def test_unknown_filename_denied(self):
  with self.assertRaises(ValueError):verify_url("https://zenodo.org/records/16918146/files/Extra.xlsx")
 def test_non_https_denied(self):
  with self.assertRaises(ValueError):verify_url("http://zenodo.org/records/16918146/files/LinearPGG_ExperimentalData.xlsx")
 def test_fake_host_denied(self):
  with self.assertRaises(ValueError):verify_url("https://zenodo.org.evil.com/records/16918146/files/LinearPGG_ExperimentalData.xlsx")
 def test_person_url_denied(self):
  with self.assertRaises(ValueError):verify_url("https://person@zenodo.org/records/16918146/files/LinearPGG_ExperimentalData.xlsx")
 def test_error_does_not_echo_subject_token(self):
  e=urllib.error.URLError("secret-subject-token999")
  self.assertNotIn("999",safe_failure(e))

class StructuredWorkbookSource(unittest.TestCase):
 def test_ten_sheet_synthetic_pass(self):
  raw=fake_book()
  d=source_summary(raw,fake_spec(raw))
  self.assertEqual(d["sheet_count"],10)
  self.assertEqual(d["author_wholefile_SHA256"],hashlib.sha256(raw).hexdigest())
 def test_no_synthetic_cell_values_emitted(self):
  raw=fake_book()
  self.assertNotIn("'42'",str(source_summary(raw,fake_spec(raw))))
 def test_record_sheet_dimensions(self):
  raw=fake_book();d=source_summary(raw,fake_spec(raw))
  self.assertEqual(d["sheet_shapes_and_boundary_nonempty_positions_only"]["FE_Type 1"]["columns_including_possible_surveys"],22)
 def test_wrong_author_md5_blocks(self):
  raw=fake_book();s=fake_spec(raw);s["original_author_md5"]="0"*32
  with self.assertRaisesRegex(ValueError,"MD5_MISMATCH"):source_summary(raw,s)
 def test_zero_source_blocks(self):
  with self.assertRaises(ValueError):source_summary(b"",fake_spec(b""))
 def test_not_xlsx_zip_blocks(self):
  raw=b"A"*300
  with self.assertRaises(ValueError):source_summary(raw,fake_spec(raw))
 def test_oversized_source_blocks(self):
  raw=b"X"*(MAX_RAW_BYTES+1)
  with self.assertRaises(ValueError):source_summary(raw,fake_spec(raw))
 def test_insufficient_original_sheet_count(self):
  b=io.BytesIO();w=Workbook();w.active.cell(11,24,1);w.save(b)
  with self.assertRaisesRegex(ValueError,"TEN_TREATMENT_SHEETS"):
   source_summary(b.getvalue(),fake_spec(b.getvalue()))

if __name__=="__main__":unittest.main()

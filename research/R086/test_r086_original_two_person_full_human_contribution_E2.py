#!/usr/bin/env python3
"""R086 synthetic only tests. No original person, 2025 workbook or real contribution."""
import hashlib,io,math,pathlib,sys,unittest
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parent))
import r086_original_two_person_full_human_contribution_E2 as research
from r086_original_two_person_full_human_contribution_E2 import (
 eligible_numeric,own_trajectories_from_four_rows,summary,analyse_one_workbook)

def make_four():
 rows=[[None]*24 for _ in range(4)]
 for i in range(4):
  rows[i][0]="synthetic-actor-0" if i==0 else None
  rows[i][1]="synthetic-pair"
  for col in range(2,24):rows[i][col]=i*100+col
 return rows

def fake_book():
 w=Workbook();w.remove(w.active)
 for arm,prefix in research.SOURCE_PREFIXES["two_person_threshold"].items():
  for role in (1,2):
   x=w.create_sheet(f"{prefix}_Type{role}")
   x.append(["PID","pair"]+["round"]*20)
   for local in range(4):
    rec=[f"synthetic-{arm}-{role}" if local==0 else None,
      f"pair-{arm}"]+list(range(local*10,local*10+20))
    x.append(rec)
 b=io.BytesIO();w.save(b)
 return b.getvalue()

class NumericAndRoleTests(unittest.TestCase):
 def test_zero_contribution_valid(self):self.assertEqual(eligible_numeric(0),0)
 def test_integer_valid(self):self.assertEqual(eligible_numeric(5),5)
 def test_float_valid(self):self.assertEqual(eligible_numeric(3.5),3.5)
 def test_booleans_are_not_contributions(self):
  with self.assertRaises(ValueError):eligible_numeric(True)
 def test_none_missing(self):
  with self.assertRaises(ValueError):eligible_numeric(None)
 def test_text_invalid(self):
  with self.assertRaises(ValueError):eligible_numeric("10")
 def test_negative_invalid(self):
  with self.assertRaises(ValueError):eligible_numeric(-1)
 def test_nan_invalid(self):
  with self.assertRaises(ValueError):eligible_numeric(float("nan"))
 def test_inf_invalid(self):
  with self.assertRaises(ValueError):eligible_numeric(float("inf"))
 def test_type1_own_role_swaps(self):
  a,b=own_trajectories_from_four_rows(make_four(),1,22,21)
  self.assertEqual(a[0],2)
  self.assertEqual(b[0],302)
  self.assertEqual([len(a),len(b)],[22,21])
 def test_type2_own_role_swaps(self):
  a,b=own_trajectories_from_four_rows(make_four(),2,22,21)
  self.assertEqual(a[0],102)
  self.assertEqual(b[0],202)
 def test_wrong_role_fails(self):
  with self.assertRaises(ValueError):own_trajectories_from_four_rows(make_four(),3,20,20)
 def test_incomplete_four_row_fails(self):
  with self.assertRaises(ValueError):own_trajectories_from_four_rows(make_four()[:3],1,20,20)
 def test_missing_round_fails(self):
  with self.assertRaises(ValueError):own_trajectories_from_four_rows(make_four(),1,100,20)

class SummaryTests(unittest.TestCase):
 def test_constant_all(self):
  s=summary([[1]*20,[2]*20])
  self.assertEqual(s["mean_human_first_last_delta"],0)
  self.assertEqual(s["focal_humans_unchanged"],2)
 def test_upward_and_downward(self):
  s=summary([[1,3],[5,2],[3,3]])
  self.assertEqual(s["focal_humans_increased"],1)
  self.assertEqual(s["focal_humans_decreased"],1)
  self.assertEqual(s["focal_humans_unchanged"],1)
  self.assertAlmostEqual(s["mean_human_first_last_delta"],-1/3)
 def test_round_means_correct(self):
  self.assertEqual(summary([[0,4],[2,6]])["mean_original_contribution_by_round"],[1,5])
 def test_no_actors_fail(self):
  with self.assertRaises(ValueError):summary([])
 def test_mismatched_rounds_fail(self):
  with self.assertRaises(ValueError):summary([[1],[1,2]])
 def test_no_ids_leak_in_summary(self):
  self.assertNotIn("person_id",str(summary([[1,2],[2,1]])))

class CompleteSyntheticUnfittedFullPipeline(unittest.TestCase):
 def pipeline(self,raw=None):
  if raw is None:raw=fake_book()
  temporary=dict(research.SOURCES["two_person_threshold"])
  temporary["original_author_md5"]=hashlib.md5(raw).hexdigest()
  temporary["published_number_of_humans_by_arm"]={x:2 for x in ("FE","EI","PI","AI","MI")}
  with patch.dict(research.SOURCES,{"two_person_threshold":temporary}):
   with patch.dict(research.PINNED_AUTHOR_SHA256,{"two_person_threshold":hashlib.sha256(raw).hexdigest()}):
    return analyse_one_workbook(raw,"two_person_threshold")
 def test_full_ten_original_sheet_role_synthetic(self):
  r=self.pipeline()
  self.assertEqual(r["source_participants_reconstructed"],10)
  self.assertEqual(len(r["original_focal_actor_treatment_session_scoring"]),5)
 def test_synthetic_session_counts_are_original_humans_not_fourfold(self):
  r=self.pipeline()
  for a in r["original_focal_actor_treatment_session_scoring"].values():
   self.assertEqual(a["session1"]["original_focal_humans"],2)
   self.assertEqual(a["session2"]["original_focal_humans"],2)
 def test_swap_has_effect_on_person_role_results(self):
  r=self.pipeline()["original_focal_actor_treatment_session_scoring"]["FE"]
  self.assertNotEqual(
   r["focal_roles"]["type1"]["session2"]["mean_first_original_contribution"],
   r["focal_roles"]["type2"]["session2"]["mean_first_original_contribution"])
 def test_no_synthetic_person_id_in_aggregate(self):
  r=self.pipeline()
  self.assertNotIn("synthetic-FE",str(r))
  self.assertNotIn("pair-FE",str(r))
 def test_sum_synthetic_2_role_groups_equals_total(self):
  r=self.pipeline()["original_focal_actor_treatment_session_scoring"]["AI"]
  self.assertEqual(r["focal_roles"]["type1"]["session1"]["original_focal_humans"]+
                   r["focal_roles"]["type2"]["session1"]["original_focal_humans"],2)
 def test_source_sha_rejected(self):
  raw=fake_book()
  with self.assertRaises(ValueError):analyse_one_workbook(raw,"two_person_threshold")
 def test_source_corrupted_workbook_rejected(self):
  with self.assertRaises(ValueError):self.pipeline(b"PK"+b"wrong"*100)
 def test_wrong_names_rejected(self):
  b=io.BytesIO();w=Workbook();w.save(b)
  with self.assertRaisesRegex(ValueError,"SHEET_NAMES"):
   self.pipeline(b.getvalue())
 def test_incomplete_actor_id_rejected(self):
  w=Workbook();w.remove(w.active)
  for arm,prefix in research.SOURCE_PREFIXES["two_person_threshold"].items():
   for role in (1,2):
    ws=w.create_sheet(f"{prefix}_Type{role}")
    ws.append(["head"]*22)
    for ix in range(4):
     ws.append([None if ix==0 else "p","g"]+[1]*20)
  b=io.BytesIO();w.save(b)
  with self.assertRaisesRegex(ValueError,"ACTOR_FIRST_ROW_ID"):self.pipeline(b.getvalue())

if __name__=="__main__":unittest.main()

#!/usr/bin/env python3
"""R086 original two-person linear+threshold PGG fixed-Zenodo MD5 and Excel schema.

Phase A: no human decision values, no ID, no actor rows ever exported.
"""
from __future__ import annotations
import argparse,hashlib,io,json,pathlib,urllib.parse,urllib.request,urllib.error
import openpyxl

ORIGINAL_RECORD=16918146
ORIGINAL_DOI="10.5281/zenodo.16918146"
MAX_RAW_BYTES=950_000
SOURCES={
 "two_person_linear":{
   "filename":"LinearPGG_ExperimentalData.xlsx",
   "original_author_md5":"fc71d20de62c259629ceb56559d8a69e",
   "published_number_of_humans_by_arm":{"FE":114,"EI":110,"PI":110,"AI":106,"MI":110},
   "author_first_session_rounds":22,"author_second_session_rounds":21},
 "two_person_threshold":{
   "filename":"ThresholdPGG_ExperimentalData.xlsx",
   "original_author_md5":"c40171abb172da8e16b0839db98a2cf2",
   "published_number_of_humans_by_arm":{"FE":110,"EI":118,"PI":112,"AI":104,"MI":110},
   "author_first_session_rounds":20,"author_second_session_rounds":20}}

def exact_url(filename):
 if filename not in {v["filename"] for v in SOURCES.values()}:
  raise ValueError("R086_UNFROZEN_AUTHOR_ORIGINAL_FILENAME")
 return "https://zenodo.org/records/16918146/files/"+urllib.parse.quote(filename)+"?download=1"

def verify_url(url):
 t=urllib.parse.urlsplit(url)
 if (t.scheme!="https" or t.hostname!="zenodo.org" or t.username
     or t.password or t.port or t.fragment):
  raise ValueError("R086_UNOFFICIAL_OR_REDIRECTED_HTTPS_LOCATION")
 if "/records/16918146/files/" not in t.path or not t.path.startswith("/records/16918146/files/"):
  raise ValueError("R086_WRONG_AUTHOR_ZENODO_RECORD")
 if urllib.parse.unquote(t.path.rsplit("/",1)[-1]) not in [x["filename"] for x in SOURCES.values()]:
  raise ValueError("R086_SOURCE_FILE_NOT_PINNED")
 if ".." in urllib.parse.unquote(t.path).split("/"):
  raise ValueError("R086_PATH_TRAVERSAL")
 return True

def source_summary(raw:bytes,spec:dict):
 if not 0<len(raw)<=MAX_RAW_BYTES:
  raise ValueError("R086_ORIGINAL_XLSX_MISSING_OR_SIZE_EXCEEDS_LIMIT")
 if hashlib.md5(raw).hexdigest()!=spec["original_author_md5"]:
  raise ValueError("R086_PUBLISHED_AUTHOR_ORIGINAL_MD5_MISMATCH")
 try:
  book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 except (ValueError,OSError,TypeError,EOFError,KeyError,zipfile.BadZipFile) as e:
  raise ValueError("R086_REAL_AUTHOR_XLSX_CANNOT_BE_READ") from e
 try:
  sheets={}
  for ws in book.worksheets:
   if ws.max_row>10000 or ws.max_column>100 or ws.max_row<10 or ws.max_column<22:
    raise ValueError("R086_AUTHOR_SHEET_SIZE_OR_ROUND_COLUMNS_INVALID")
   first=[i+1 for i,v in enumerate(next(ws.iter_rows(min_row=1,max_row=1,values_only=True))) if v is not None]
   last=[i+1 for i,v in enumerate(next(ws.iter_rows(min_row=ws.max_row,max_row=ws.max_row,values_only=True))) if v is not None]
   sheets[ws.title]={"rows_including_possible_headers":ws.max_row,
     "columns_including_possible_surveys":ws.max_column,
     "first_row_nonempty_column_positions":first,
     "last_row_nonempty_column_positions":last}
  if len(sheets)!=10:raise ValueError("R086_AUTHOR_ORIGINAL_NOT_TEN_TREATMENT_SHEETS")
 finally:book.close()
 return {"official_publisher_wholefile_MD5_exact_match":True,
   "author_original_binary_bytes":len(raw),
   "author_wholefile_SHA256":hashlib.sha256(raw).hexdigest(),
   "workbook_valid_openpyxl":True,"sheet_count":len(sheets),
   "sheet_shapes_and_boundary_nonempty_positions_only":sheets,
   "original_author_reported_humans_by_condition_not_own_calculation":spec["published_number_of_humans_by_arm"],
   "original_author_reported_session_round_counts":{
     "first":spec["author_first_session_rounds"],
     "second":spec["author_second_session_rounds"]}}

def safe_failure(e):
 if isinstance(e,urllib.error.HTTPError):return f"OFFICIAL_ZENODO_HTTP_{e.code}"
 if isinstance(e,urllib.error.URLError):return "NETWORK_NOT_REACHABLE"
 if isinstance(e,(OSError,TimeoutError)):return "SOURCE_NETWORK_IO_ERROR"
 if isinstance(e,ValueError) and str(e).startswith("R086_"):return str(e)[:90]
 return "UNKNOWN_SOURCE_FAILURE"

# zipfile exception in openpyxl parser is expected but not rendered to output.
import zipfile

def evaluate():
 checks={}
 for key,spec in SOURCES.items():
  result={"official_author_filename":spec["filename"],
    "official_published_MD5":spec["original_author_md5"]}
  try:
   url=exact_url(spec["filename"])
   verify_url(url)
   req=urllib.request.Request(url,headers={"User-Agent":"Psychohistory-R086-publisher-source-original-XLSX-identity/1.0"})
   with urllib.request.urlopen(req,timeout=40) as response:
    verify_url(response.geturl())
    raw=response.read(MAX_RAW_BYTES+1)
   result.update(source_summary(raw,spec))
   result["source_gate"]="AUTHOR_ORIGINAL_EXCEL_MD5_AND_SHEET_SCHEMAS_VERIFIED"
  except (ValueError,urllib.error.HTTPError,urllib.error.URLError,TimeoutError,OSError,TypeError) as e:
   result["source_gate"]="FAIL_CLOSED"
   result["source_error_class"]=safe_failure(e)
  checks[key]=result
 overall=all(z["source_gate"]=="AUTHOR_ORIGINAL_EXCEL_MD5_AND_SHEET_SCHEMAS_VERIFIED" for z in checks.values())
 return {"research_id":"R086",
  "phase":"A_2025_TWO_PERSON_LINEAR_THRESHOLD_PGG_ORIGINAL_XLSX_SOURCE_AND_SCHEMA_ONLY",
  "status":"ORIGINAL_BOTH_XLSX_MD5_AND_WORKBOOK_SCHEMAS_VERIFIED" if overall else "ORIGINAL_XLSX_SOURCE_OR_SCHEMA_FAIL_CLOSED",
  "official_record":ORIGINAL_RECORD,"official_DOI":ORIGINAL_DOI,
  "source_gate_outputs":checks,
  "any_real_human_behavior_values_exported":False,
  "original_behavioural_cooperation_effects_computed":0,
  "real_human_outcome_E2":"NOT_YET_EXAMINED",
  "PESEN_v0.1":"FROZEN"}

def main():
 p=argparse.ArgumentParser();p.add_argument("--output",type=pathlib.Path,required=True)
 a=p.parse_args()
 r=evaluate()
 a.output.parent.mkdir(parents=True,exist_ok=True)
 a.output.write_text(json.dumps(r,sort_keys=True,indent=2)+"\n")
 print(json.dumps(r,sort_keys=True))
if __name__=="__main__":main()

#!/usr/bin/env python3
"""R086 exact author-publisher two-person five-arm PGG cooperation dynamics E2.

Four rows per focal human -> two own session trajectories. Never release actor IDs.
"""
from __future__ import annotations
import argparse,hashlib,io,json,math,pathlib,urllib.request,urllib.error
from collections import defaultdict
from numbers import Real
import openpyxl
from r086_original_zenodo_two_player_xlsx_source_gate import SOURCES,exact_url,verify_url,MAX_RAW_BYTES

PINNED_AUTHOR_SHA256={
 "two_person_linear":"b423cda39bfd7d84f2edb7f56b4bccc0d03164743bb4df851d20f049725fb507",
 "two_person_threshold":"825b1cd3226aeca7a19d76d9594f655114b12059648f9ca03b3291c3cf9ab200"}
SOURCE_PREFIXES={
 "two_person_linear":{
  "FE":"Full_equlity","EI":"Endowment_inequality","PI":"Productivity_inequality",
  "AI":"Aligned_inequality","MI":"Misaliged_inequality"},
 "two_person_threshold":{
  "FE":"Full_equality","EI":"Endowment_inequality","PI":"Productivity_inequality",
  "AI":"Aligned_inequality","MI":"Misaligned_inequality"}}
PHASE_B_PROTOCOL="protocols/R086_PHASE_B_TWO_PERSON_FULL_ORIGINAL_BEHAVIOUR_WITHIN_SESSION_FROZEN_SCORE_2026-10-17.md"

def eligible_numeric(x):
 if isinstance(x,bool) or not isinstance(x,Real) or not math.isfinite(float(x)) or float(x)<0:
  raise ValueError("R086_MISSING_NONNUMERIC_NEGATIVE_OR_NONFINITE_TRUE_HUMAN_CONTRIBUTION")
 return float(x)

def own_trajectories_from_four_rows(rows,focal_role,n_first,n_second):
 if len(rows)!=4 or focal_role not in (1,2) or n_first<1 or n_second<1:
  raise ValueError("R086_FOUR_ROW_PERSON_OR_ROLE_OR_DURATION_INVALID")
 own_a=rows[0] if focal_role==1 else rows[1]
 own_b=rows[3] if focal_role==1 else rows[2]
 if len(own_a)<2+n_first or len(own_b)<2+n_second:
  raise ValueError("R086_EXPECTED_C_TO_LATEST_ROUND_MISSING")
 a=[eligible_numeric(own_a[k]) for k in range(2,2+n_first)]
 b=[eligible_numeric(own_b[k]) for k in range(2,2+n_second)]
 return a,b

def summary(actors):
 if not actors:raise ValueError("R086_EMPTY_ORIGINAL_FOCAL_ROLE_RISK_SET")
 n=len(actors);length=len(actors[0])
 if any(len(t)!=length for t in actors):
  raise ValueError("R086_INCONSISTENT_SOURCE_SESSION_ROUNDS")
 delta=[x[-1]-x[0] for x in actors]
 return {"original_focal_humans":n,
  "source_session_rounds":length,
  "mean_original_contribution_by_round":[sum(x[i] for x in actors)/n for i in range(length)],
  "mean_first_original_contribution":sum(x[0] for x in actors)/n,
  "mean_final_original_contribution":sum(x[-1] for x in actors)/n,
  "mean_human_first_last_delta":sum(delta)/n,
  "focal_humans_increased":sum(z>0 for z in delta),
  "focal_humans_decreased":sum(z<0 for z in delta),
  "focal_humans_unchanged":sum(z==0 for z in delta),
  "no_independent_human_group_or_causal_pvalue_assigned":True}

def analyse_one_workbook(raw,game):
 spec=SOURCES[game]
 if len(raw)>MAX_RAW_BYTES or not raw:raise ValueError("R086_REAL_SOURCE_FILE_SIZE_REJECTED")
 if hashlib.md5(raw).hexdigest()!=spec["original_author_md5"]:
  raise ValueError("R086_REAL_AUTHOR_MD5_NOT_MATCHED")
 if hashlib.sha256(raw).hexdigest()!=PINNED_AUTHOR_SHA256[game]:
  raise ValueError("R086_PRE_BEHAVIOR_SOURCE_SHA256_PARITY_FAILED")
 try:book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 except (OSError,ValueError,KeyError,EOFError) as e:
  raise ValueError("R086_ORIGINAL_XLSX_REOPEN_NOT_POSSIBLE") from e
 prefix=SOURCE_PREFIXES[game]
 expected_sheets={f"{prefix[arm]}_Type{role}" for arm in prefix for role in (1,2)}
 if set(book.sheetnames)!=expected_sheets:
  book.close()
  raise ValueError("R086_FROZEN_ORIGINAL_AUTHOR_TEN_SHEET_NAMES_MISMATCH")
 arm_scores={}
 for arm in sorted(prefix):
  n_author=spec["published_number_of_humans_by_arm"][arm]
  all_by_session={1:[],2:[]}
  roles_by_session={1:{1:[],2:[]},2:{1:[],2:[]}}
  for role in (1,2):
   ws=book[f"{prefix[arm]}_Type{role}"]
   if ws.max_row!=2*n_author+1:
    book.close()
    raise ValueError("R086_ORIGINAL_HUMAN_ROSTER_SHEET_ROW_COUNT_MISMATCH")
   rows=list(ws.iter_rows(min_row=2,max_row=ws.max_row,values_only=True))
   if len(rows)!=2*n_author or len(rows)%4:
    book.close()
    raise ValueError("R086_AUTHOR_FOUR_ROW_ACTOR_UNIVERSE_NOT_EXHAUSTIVE")
   ids=set()
   for offset in range(0,len(rows),4):
    four=rows[offset:offset+4]
    actorid=four[0][0]
    if actorid is None or isinstance(actorid,bool) or not isinstance(actorid,(str,int,float)):
     book.close()
     raise ValueError("R086_AUTHOR_ORIGINAL_ACTOR_FIRST_ROW_ID_MISSING_OR_MALFORMED")
    identity=str(actorid)
    if identity in ids:
     book.close()
     raise ValueError("R086_AUTHOR_ORIGINAL_DUPLICATE_FOCAL_ID_WITHIN_SHEET")
    ids.add(identity)
    a,b=own_trajectories_from_four_rows(four,role,
      spec["author_first_session_rounds"],spec["author_second_session_rounds"])
    roles_by_session[1][role].append(a)
    roles_by_session[2][role].append(b)
    all_by_session[1].append(a);all_by_session[2].append(b)
   if len(ids)!=n_author//2:
    book.close()
    raise ValueError("R086_PUBLISHED_ORIGINAL_AUTHOR_ROLE_ROSTER_NOT_MET")
  s1=summary(all_by_session[1]);s2=summary(all_by_session[2])
  if s1["original_focal_humans"]!=n_author or s2["original_focal_humans"]!=n_author:
   book.close()
   raise ValueError("R086_ORIGINAL_AUTHOR_ACTOR_DENOMINATOR_CHANGED")
  arm_scores[arm]={
   "author_focal_humans_per_session":n_author,
   "session1":s1,"session2":s2,
   "focal_roles":{
      "type1":{"session1":summary(roles_by_session[1][1]),"session2":summary(roles_by_session[2][1])},
      "type2":{"session1":summary(roles_by_session[1][2]),"session2":summary(roles_by_session[2][2])}},
   "same_arm_both_original_sessions_delta_sign_consistent":
     (s1["mean_human_first_last_delta"]>0 and s2["mean_human_first_last_delta"]>0)
     or (s1["mean_human_first_last_delta"]<0 and s2["mean_human_first_last_delta"]<0)
     or (s1["mean_human_first_last_delta"]==0 and s2["mean_human_first_last_delta"]==0)}
 book.close()
 if sum(x["author_focal_humans_per_session"] for x in arm_scores.values())!=sum(spec["published_number_of_humans_by_arm"].values()):
  raise ValueError("R086_FIVE_TREATMENT_TOTAL_ORIGINAL_AUTHORS_NOT_MATCHED")
 return {"author_original_file":spec["filename"],
   "author_publisher_MD5_strict_match":True,
   "pinned_first_source_SHA256_match":True,
   "whole_original_source_SHA256":PINNED_AUTHOR_SHA256[game],
   "source_participants_reconstructed":sum(spec["published_number_of_humans_by_arm"].values()),
   "participant_units":"distinct_original_focus_actors_not_partner_duplicate_rows",
   "sessions_per_actor":2,
   "original_focal_actor_treatment_session_scoring":arm_scores}

def true_author_result():
 games={}
 for name,spec in SOURCES.items():
  url=exact_url(spec["filename"]);verify_url(url)
  req=urllib.request.Request(url,headers={
   "User-Agent":"PsychologicalHistory-R086-original-human-source-5arm-2025/1.0"})
  with urllib.request.urlopen(req,timeout=45) as res:
   verify_url(res.geturl())
   raw=res.read(MAX_RAW_BYTES+1)
  games[name]=analyse_one_workbook(raw,name)
 return {"research_id":"R086",
  "status":"R086_FULL_AUTHOR_ORIGINAL_TWO_GAME_FIVE_TREATMENT_HUMAN_BEHAVIOR_E2_EXECUTED",
  "real_human_scientific_grade":"E2_RETROSPECTIVE_AUTHOR_ORIGINAL_XLSX_CONTRIBUTIONS",
  "source_DOI":"10.5281/zenodo.16918146",
  "frozen_score_protocol":PHASE_B_PROTOCOL,
  "author_original_two_games":games,
  "original_total_distinct_human_focus_actors_across_two_games":sum(x["source_participants_reconstructed"] for x in games.values()),
  "original_two_sessions_same_actors_not_independent_new_cohorts":True,
  "actual_original_behavioral_values_computed":True,
  "individual_human_IDs_exported":False,
  "raw_human_behaviour_rows_exported":False,
  "original_XLSX_files_in_git":False,
  "across_treatment_randomized_causal_effect":"NOT_IDENTIFIED",
  "long_term_historical_dynamics":"NOT_VALIDATED",
  "PESEN_v0.1":"FROZEN"}

def main():
 parser=argparse.ArgumentParser();parser.add_argument("--output",type=pathlib.Path,required=True)
 a=parser.parse_args()
 try:r=true_author_result()
 except (ValueError,OSError,EOFError,TypeError,urllib.error.HTTPError,urllib.error.URLError) as e:
  reason=str(e)[:90] if isinstance(e,ValueError) and str(e).startswith("R086_") else "SOURCE_IO_OR_SCHEMA_FAILURE"
  r={"research_id":"R086","status":"R086_TRUE_AUTHOR_ORIGINAL_HUMAN_SCORING_STRICT_FAIL_CLOSED",
     "safe_error_class":reason,"original_human_results_exported":0,"PESEN_v0.1":"FROZEN"}
 a.output.parent.mkdir(parents=True,exist_ok=True)
 a.output.write_text(json.dumps(r,indent=2,sort_keys=True)+"\n")
 print(json.dumps({"status":r["status"],"source_real_scores":{
  k:{"N":v["source_participants_reconstructed"],
     "delta_by_arm_and_session":{a:{q:x[q]["mean_human_first_last_delta"] for q in ("session1","session2")}
                                    for a,x in v["original_focal_actor_treatment_session_scoring"].items()}}
  for k,v in r.get("author_original_two_games",{}).items()},
  "safe_error_class":r.get("safe_error_class")},sort_keys=True))
 if r["status"]!="R086_FULL_AUTHOR_ORIGINAL_TWO_GAME_FIVE_TREATMENT_HUMAN_BEHAVIOR_E2_EXECUTED":
  raise RuntimeError("R086_ORIGINAL_HUMAN_E2_STRICT_SCORING_GATE_FAILED")
if __name__=="__main__":main()
